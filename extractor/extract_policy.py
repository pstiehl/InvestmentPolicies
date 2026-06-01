#!/usr/bin/env python3
"""
Extract a single client investment policy file → canonical JSON.

Inputs:  .docx or .xlsx
Output:  data/extracted/<client_slug>.json conforming to schema/policy.schema.json

Privacy rules (mandatory):
  - Client name is replaced with a stable pseudonym ("Client_<hash6>") before
    being sent to the LLM, and re-attached after the response.
  - Specific dollar amounts > $1,000 are redacted to "$<REDACTED>" before LLM
    call, then matched back into the response using a token map.
  - No prompt bodies are persisted. Only metadata (timestamps, token counts,
    cost, model, prompt hash) lands in data/llm_audit.jsonl.

Usage:
  python extractor/extract_policy.py <path-to-policy-file> \
      --client "Merck Sharp & Dohme" \
      --policy-year 2026 \
      --out data/extracted/merck_sharp_and_dohme.json
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import time
from datetime import datetime, timezone
from pathlib import Path

EXTRACTOR_VERSION = "0.1.0"
MODEL = os.environ.get("ALMP_MODEL", "claude-opus-4-7")

REPO_ROOT = Path(__file__).resolve().parent.parent
SCHEMA_PATH = REPO_ROOT / "schema" / "policy.schema.json"
AUDIT_PATH = REPO_ROOT / "data" / "llm_audit.jsonl"
EXTRACTED_DIR = REPO_ROOT / "data" / "extracted"


# ---------- file → plain text ----------

def read_docx(path: Path) -> str:
    from docx import Document
    doc = Document(str(path))
    out = []
    for p in doc.paragraphs:
        if p.text.strip():
            out.append(p.text)
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def read_xlsx(path: Path) -> str:
    from openpyxl import load_workbook
    wb = load_workbook(str(path), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        out.append(f"## Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                out.append(" | ".join(cells))
    return "\n".join(out)


def read_policy_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return read_docx(path)
    if suffix in (".xlsx", ".xlsm"):
        return read_xlsx(path)
    raise ValueError(f"Unsupported file type: {suffix}. Expected .docx or .xlsx.")


# ---------- redaction ----------

DOLLAR_RE = re.compile(r"\$\s?([0-9]{1,3}(?:,[0-9]{3})+|[0-9]{4,})(?:\.\d+)?")


def redact_for_llm(text: str, client_name: str) -> tuple[str, dict]:
    """Returns (redacted_text, token_map).

    - Client name → stable pseudonym
    - $X where X > 1000 → "$<DOLLAR_n>" token
    """
    pseudonym = f"Client_{hashlib.sha256(client_name.encode()).hexdigest()[:6]}"
    token_map = {"client_pseudonym": pseudonym, "dollars": {}}

    redacted = re.sub(
        re.escape(client_name),
        pseudonym,
        text,
        flags=re.IGNORECASE,
    )

    counter = [0]
    def _dollar_sub(m):
        raw = m.group(0)
        try:
            numeric = float(m.group(1).replace(",", ""))
        except ValueError:
            return raw
        if numeric <= 1000:
            return raw
        counter[0] += 1
        tok = f"$<DOLLAR_{counter[0]}>"
        token_map["dollars"][tok] = raw
        return tok

    redacted = DOLLAR_RE.sub(_dollar_sub, redacted)
    return redacted, token_map


def restore(text_or_obj, token_map: dict, client_name: str):
    """Reverse the redaction on any string fields recursively."""
    pseudonym = token_map["client_pseudonym"]
    dollars = token_map["dollars"]

    def restore_string(s: str) -> str:
        s = s.replace(pseudonym, client_name)
        for tok, original in dollars.items():
            s = s.replace(tok, original)
        return s

    def walk(node):
        if isinstance(node, str):
            return restore_string(node)
        if isinstance(node, list):
            return [walk(x) for x in node]
        if isinstance(node, dict):
            return {k: walk(v) for k, v in node.items()}
        return node

    return walk(text_or_obj)


# ---------- LLM ----------

SYSTEM_PROMPT = """You are an investment compliance analyst at ALM First.

You are given the text of a single client investment policy. Your job is to
extract every concrete trading limit and restriction into a structured JSON
object that conforms exactly to the schema below.

Rules:
- Output JSON ONLY. No prose, no markdown fences, no commentary.
- For every asset class the policy explicitly addresses, emit one entry in
  `asset_classes`. Use the canonical category names from `x-canonical-categories`
  when possible; if the policy describes a sub-structure (e.g. "Agency CMO PAC"),
  use `subcategory` to capture it.
- If a limit is unspecified, set the field to null. Do NOT invent limits.
- Convert maturity expressions to years (decimal allowed): "18 months" → 1.5.
- Convert percentages to the numeric percent value (15 not 0.15).
- For `single_security_limit`, preserve the original expression as a string
  ("10% of Net Worth" or "$250,000") because units vary.
- For `min_credit_rating`, normalize to S&P scale (AA-, A, BBB+, etc.).
- If the policy explicitly prohibits a structure (e.g. "support tranche CMOs
  are prohibited"), add it to `structures_prohibited` for the relevant class.
- If you encounter a clause that constrains trading but doesn't fit the schema,
  copy it verbatim into `unmapped_clauses` rather than forcing it.
- Set `extraction_confidence` between 0 and 1 honestly: 1.0 means every clause
  mapped cleanly; lower means some clauses were ambiguous or unmapped.
- The client name has been replaced with a pseudonym; preserve it as-is in your
  output (post-processing will swap it back).
- Dollar amounts above $1,000 have been replaced with tokens like
  $<DOLLAR_3>. Preserve those tokens verbatim where you see them.
"""

USER_TEMPLATE = """SCHEMA:
{schema}

POLICY TEXT (client: {pseudonym}, source file: {filename}):
---
{policy_text}
---

Emit the JSON object now."""


def call_llm(redacted_text: str, pseudonym: str, filename: str, schema: dict) -> tuple[dict, dict]:
    """Returns (parsed_json, audit_metadata)."""
    try:
        import anthropic
    except ImportError:
        sys.exit("pip install anthropic")

    client = anthropic.Anthropic()
    user_prompt = USER_TEMPLATE.format(
        schema=json.dumps(schema, indent=2),
        pseudonym=pseudonym,
        filename=filename,
        policy_text=redacted_text,
    )

    t0 = time.time()
    resp = client.messages.create(
        model=MODEL,
        max_tokens=8000,
        system=SYSTEM_PROMPT,
        messages=[{"role": "user", "content": user_prompt}],
    )
    dt = time.time() - t0

    raw = "".join(b.text for b in resp.content if hasattr(b, "text")).strip()
    if raw.startswith("```"):
        # tolerate accidental fenced JSON
        raw = re.sub(r"^```(?:json)?\s*", "", raw)
        raw = re.sub(r"\s*```$", "", raw)
    parsed = json.loads(raw)

    audit = {
        "ts": datetime.now(timezone.utc).isoformat(),
        "model": MODEL,
        "prompt_hash": hashlib.sha256(user_prompt.encode()).hexdigest(),
        "elapsed_s": round(dt, 2),
        "input_tokens": resp.usage.input_tokens,
        "output_tokens": resp.usage.output_tokens,
        # cost figures intentionally not hardcoded — compute downstream
    }
    return parsed, audit


# ---------- main ----------

def slugify(name: str) -> str:
    s = re.sub(r"[^a-zA-Z0-9]+", "_", name).strip("_").lower()
    return s or "client"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("policy_path", type=Path)
    ap.add_argument("--client", required=True, help="Client display name as in S:\\Clients\\<ClientName>")
    ap.add_argument("--policy-year", default=None, help="Year folder under policies/, e.g. 2026")
    ap.add_argument("--policy-date", default=None, help="Effective date YYYY-MM-DD if known")
    ap.add_argument("--out", type=Path, default=None)
    ap.add_argument("--dry-run", action="store_true", help="Print redacted text instead of calling the API")
    args = ap.parse_args()

    if not args.policy_path.exists():
        sys.exit(f"Not found: {args.policy_path}")

    EXTRACTED_DIR.mkdir(parents=True, exist_ok=True)
    AUDIT_PATH.parent.mkdir(parents=True, exist_ok=True)

    schema = json.loads(SCHEMA_PATH.read_text())

    text = read_policy_text(args.policy_path)
    file_hash = hashlib.sha256(args.policy_path.read_bytes()).hexdigest()

    redacted, token_map = redact_for_llm(text, args.client)

    if args.dry_run:
        print(f"--- REDACTED ({len(redacted)} chars) ---")
        print(redacted[:4000])
        print("--- TOKEN MAP ---")
        print(json.dumps({"pseudonym": token_map["client_pseudonym"],
                         "dollar_tokens": len(token_map["dollars"])}, indent=2))
        return

    print(f"    [llm] sending {len(redacted)} chars to {MODEL}…", file=sys.stderr, flush=True)
    parsed, audit = call_llm(redacted, token_map["client_pseudonym"], args.policy_path.name, schema)
    print(f"    [llm] response in {audit['elapsed_s']}s, {audit['input_tokens']} in / {audit['output_tokens']} out tokens", file=sys.stderr, flush=True)

    # Re-attach real client name + real dollar amounts
    parsed = restore(parsed, token_map, args.client)

    # Server-side metadata we set authoritatively
    parsed["client"] = args.client
    parsed["policy_file"] = args.policy_path.name
    parsed["policy_file_hash"] = file_hash
    parsed["extracted_at"] = datetime.now(timezone.utc).isoformat()
    parsed["extractor_version"] = EXTRACTOR_VERSION
    if args.policy_year:
        parsed["policy_year_folder"] = args.policy_year
    if args.policy_date:
        parsed["policy_date"] = args.policy_date

    # Validate
    try:
        import jsonschema
        jsonschema.validate(parsed, schema)
    except ImportError:
        print("[warn] jsonschema not installed; skipping validation", file=sys.stderr)
    except Exception as e:
        print(f"[warn] schema validation failed: {e}", file=sys.stderr)

    out_path = args.out or EXTRACTED_DIR / f"{slugify(args.client)}.json"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(parsed, indent=2))
    print(f"[ok] wrote {out_path}", file=sys.stderr)

    with AUDIT_PATH.open("a") as f:
        f.write(json.dumps({**audit,
                            "client_pseudonym": token_map["client_pseudonym"],
                            "out_path": str(out_path),
                            "input_file_hash": file_hash}) + "\n")


if __name__ == "__main__":
    main()
